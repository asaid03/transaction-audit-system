import io
import unittest

import openpyxl
import pandas as pd

from transaction_audit.audit import (
    build_audit_summary,
    build_audit_workbook,
    build_bundle_audit_workbook,
    issue_counts_dataframe,
    review_findings_dataframe,
)
from transaction_audit.bundle_workflow import run_bundle_audit
from transaction_audit.datasets import AuditBundle, AuditDataset, DatasetType
from transaction_audit.schema import apply_schema
from transaction_audit.validation import validate_transactions


class AuditTests(unittest.TestCase):
    def test_issue_counts_group_by_type_and_severity(self):
        dataframe = pd.DataFrame(
            {
                "transaction_id": ["TXN-001", "TXN-002", "TXN-002"],
                "transaction_date": ["2026-04-01", "not-a-date", "2026-04-02"],
                "account_id": ["ACC-1", "ACC-1", "ACC-1"],
                "amount": ["0", "10", "10"],
                "currency": ["GBP", "GBP", "GBP"],
                "counterparty": ["Vendor A", "Vendor B", "Vendor B"],
            }
        )

        result = validate_transactions(dataframe)
        issue_counts = issue_counts_dataframe(result)

        self.assertIn("issue_type", issue_counts.columns)
        self.assertIn("severity", issue_counts.columns)
        self.assertIn("count", issue_counts.columns)
        self.assertIn("duplicate_transaction_id", issue_counts["issue_type"].tolist())

    def test_build_audit_workbook_contains_review_sheets(self):
        dataframe = pd.DataFrame(
            {
                "transaction_id": ["TXN-001"],
                "transaction_date": ["2026-04-01"],
                "account_id": ["ACC-1"],
                "amount": ["10"],
                "currency": ["GBP"],
                "counterparty": ["Vendor A"],
            }
        )
        schema_result = apply_schema(dataframe)
        result = validate_transactions(schema_result.transactions)
        summary = build_audit_summary(result, "sample.csv")

        workbook_bytes = build_audit_workbook(result, schema_result, summary)
        workbook = openpyxl.load_workbook(io.BytesIO(workbook_bytes), read_only=True)

        self.assertEqual(
            {
                "summary",
                "review_findings",
                "finding_summary",
                "raw_issues",
                "issue_counts",
                "schema_mapping",
                "transactions",
            },
            set(workbook.sheetnames),
        )

    def test_review_findings_explain_issue_and_action(self):
        dataframe = pd.DataFrame(
            {
                "transaction_id": ["TXN-001"],
                "transaction_date": ["2026-04-01"],
                "account_id": ["ACC-1"],
                "amount": ["0"],
                "currency": ["GBP"],
                "counterparty": ["Vendor A"],
            }
        )
        result = validate_transactions(dataframe)

        findings = review_findings_dataframe(result.issues_dataframe())

        self.assertIn("finding", findings.columns)
        self.assertIn("recommended_action", findings.columns)
        self.assertIn("review_status", findings.columns)
        self.assertEqual("Needs review", findings.loc[0, "review_status"])
        self.assertEqual("Zero amount transaction", findings.loc[0, "finding"])

    def test_review_findings_sort_high_priority_before_warnings(self):
        issues = pd.DataFrame(
            [
                {
                    "dataset": "transactions",
                    "row_number": 2,
                    "transaction_id": "TXN-001",
                    "field": "amount",
                    "issue_type": "zero_amount",
                    "severity": "warning",
                    "related_dataset": None,
                    "related_record_id": None,
                    "message": "Zero amount.",
                },
                {
                    "dataset": "transactions",
                    "row_number": 3,
                    "transaction_id": "TXN-002",
                    "field": "vendor_id",
                    "issue_type": "unknown_vendor_reference",
                    "severity": "error",
                    "related_dataset": "vendors",
                    "related_record_id": "VEN-404",
                    "message": "Vendor reference was not found.",
                },
            ]
        )

        findings = review_findings_dataframe(issues)

        self.assertEqual("High", findings.loc[0, "priority"])
        self.assertEqual("Payment references an unapproved or missing vendor", findings.loc[0, "finding"])

    def test_build_bundle_audit_workbook_contains_multi_file_review_sheets(self):
        transactions = pd.DataFrame(
            {
                "transaction_id": ["TXN-001", "TXN-002"],
                "transaction_date": ["2026-05-01", "2026-05-02"],
                "account_id": ["OPS", "OPS"],
                "amount": ["100", "200"],
                "currency": ["USD", "USD"],
                "counterparty": ["Known Vendor", "Missing Vendor"],
                "vendor_id": ["VEN-001", "VEN-404"],
            }
        )
        vendors = pd.DataFrame(
            {
                "vendor_id": ["VEN-001"],
                "vendor_name": ["Known Vendor"],
            }
        )
        transaction_schema = apply_schema(transactions)
        vendor_schema = apply_schema(
            vendors,
            source_to_canonical={"vendor_id": "vendor_id", "vendor_name": "vendor_name"},
            required_columns=("vendor_id", "vendor_name"),
            allowed_columns=("vendor_id", "vendor_name"),
        )
        bundle = AuditBundle(
            transactions=AuditDataset(DatasetType.TRANSACTIONS, transaction_schema.transactions, "transactions.csv"),
            vendors=AuditDataset(DatasetType.VENDORS, vendor_schema.transactions, "vendors.csv"),
        )
        bundle_result = run_bundle_audit(bundle)

        workbook_bytes = build_bundle_audit_workbook(
            bundle_result,
            {
                DatasetType.TRANSACTIONS: transaction_schema,
                DatasetType.VENDORS: vendor_schema,
            },
        )
        workbook = openpyxl.load_workbook(io.BytesIO(workbook_bytes), read_only=True)

        self.assertEqual(
            {
                "read_me",
                "summary",
                "review_findings",
                "finding_summary",
                "dataset_summary",
                "issue_counts",
                "transaction_issues",
                "reference_issues",
                "raw_issues",
                "skipped_checks",
                "schema_mappings",
                "transactions",
                "vendors",
            },
            set(workbook.sheetnames),
        )
        self.assertEqual("unknown_vendor_reference", workbook["reference_issues"]["E2"].value)
